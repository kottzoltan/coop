#!/usr/bin/env python3
"""Munkaerő SZSZ Excel → SQL seed (partner, projekt, szerződés, munkatárs)."""

from __future__ import annotations

import json
import re
from datetime import datetime, date
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
IMPORT_DIR = ROOT / "data" / "imports" / "munkaero-2026-08-31"
OUT_SQL = ROOT / "scripts" / "sql" / "munkaero_szsz_seed.sql"
OUT_JSON = ROOT / "data" / "imports" / "munkaero-2026-08-31" / "summary.json"

IRODA = "Munkaerő Humánszolgáltató Szociális Szövetkezet - Iroda"
CEG_KOD = "munkaero"


def sql_str(v) -> str:
    if v is None:
        return "NULL"
    s = str(v).strip()
    if s == "":
        return "NULL"
    return "'" + s.replace("'", "''") + "'"


def sql_json(obj) -> str:
    return sql_str(json.dumps(obj, ensure_ascii=False))


def parse_date(v):
    if v is None or v == "":
        return None
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    s = str(v).strip()
    for fmt in ("%Y.%m.%d", "%Y-%m-%d", "%Y/%m/%d"):
        try:
            return datetime.strptime(s[:10].replace(" ", ""), fmt).date().isoformat()
        except ValueError:
            continue
    # 2023.04.22 style already
    m = re.match(r"(\d{4})[.\-/](\d{1,2})[.\-/](\d{1,2})", s)
    if m:
        return f"{int(m.group(1)):04d}-{int(m.group(2)):02d}-{int(m.group(3)):02d}"
    return None


def norm_status(s) -> str:
    if not s:
        return "aktív"
    t = str(s).strip().lower()
    if t in ("aktív", "aktiv", "active"):
        return "aktív"
    if t in ("inaktív", "inaktiv", "inactive"):
        return "inaktív"
    return str(s).strip()


def find_xlsx(pattern: str) -> Path:
    for p in IMPORT_DIR.glob("*.xlsx"):
        if p.name.startswith("~$"):
            continue
        if pattern in p.name.lower():
            return p
    raise FileNotFoundError(pattern)


def load_partners(path: Path):
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    header = [str(h).strip() if h else "" for h in rows[0]]
    out = []
    for r in rows[1:]:
        if not r or not r[0]:
            continue
        d = dict(zip(header, r))
        out.append(
            {
                "nev": str(d.get("Név") or "").strip(),
                "adoszam": (str(d.get("Adószám / Magánszemély adóazonosító jel") or "").strip() or None),
                "cegjegyzekszam": (str(d.get("Cégjegyzékszám") or "").strip() or None),
                "bankszamlaszam": (str(d.get("Bankszámlaszám") or "").strip() or None),
                "afakoros": str(d.get("Áfakörös") or "").strip().lower() in ("igen", "true", "1", "x"),
                "cim": (str(d.get("Székhely cím") or "").strip() or None),
                "kulso_id": d.get("Azonosító"),
                "agazat": (str(d.get("Nemzetgazdagási ágazat") or d.get("Nemzetgazdasági ágazat") or "").strip() or None),
            }
        )
    return out


def load_projektek(path: Path):
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    header = [str(h).strip() if h else "" for h in rows[0]]
    out = []
    for r in rows[1:]:
        if not r or not r[0]:
            continue
        d = dict(zip(header, r))
        out.append(d)
    return out


def load_jogosultsag(path: Path):
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    ids_row, names, emails, roles = rows[0], rows[1], rows[2], rows[3]
    users = []
    for col in range(2, len(ids_row)):
        raw_id = ids_row[col]
        email = emails[col]
        if not email:
            continue
        kulso = None
        if raw_id and "ID:" in str(raw_id):
            try:
                kulso = int(str(raw_id).split(":")[-1].strip())
            except ValueError:
                kulso = None
        perms = {}
        for r in rows[4:]:
            mod, action = r[0], r[1]
            if not mod or not action:
                continue
            val = r[col] if col < len(r) else None
            granted = str(val).strip().lower() in ("x", "1", "true", "igen")
            perms.setdefault(str(mod).strip(), {})[str(action).strip()] = granted
        users.append(
            {
                "kulso_id": kulso,
                "nev": str(names[col] or "").strip(),
                "email": str(email).strip().lower(),
                "szerepkorok": str(roles[col] or "").strip() or None,
                "jogosultsag": perms,
            }
        )
    return users


def main():
    partner_path = find_xlsx("partner-export")
    projekt_path = find_xlsx("projekt")
    jog_path = find_xlsx("jogosultsag")

    partners = load_partners(partner_path)
    projektek = load_projektek(projekt_path)
    users = load_jogosultsag(jog_path)

    lines = [
        "-- Auto-generated: Munkaerő SZSZ seed from Excel 2026-08-31",
        "-- Idempotent for ceg=munkaero (delete+reinsert tenant data)",
        "BEGIN;",
        "",
        f"-- Ensure tenant",
        f"INSERT INTO ceg (kod, nev) VALUES ({sql_str(CEG_KOD)}, {sql_str('Munkaerő Humánszolgáltató Szociális Szövetkezet')}) ON CONFLICT (kod) DO UPDATE SET nev = EXCLUDED.nev;",
        "SELECT id AS ceg_id FROM ceg WHERE kod = 'munkaero' \\gset",
        "",
        "-- Clear previous import for this tenant (projektek first via partner link / ceg_id)",
        "DELETE FROM partner_szerzodes WHERE partner_id IN (SELECT id FROM partner WHERE ceg_id = (SELECT id FROM ceg WHERE kod = 'munkaero'));",
        "DELETE FROM partner_kapcsolattarto WHERE partner_id IN (SELECT id FROM partner WHERE ceg_id = (SELECT id FROM ceg WHERE kod = 'munkaero'));",
        "DELETE FROM projekt WHERE ceg_id = (SELECT id FROM ceg WHERE kod = 'munkaero');",
        "DELETE FROM partner WHERE ceg_id = (SELECT id FROM ceg WHERE kod = 'munkaero');",
        "DELETE FROM munkatars WHERE ceg_id = (SELECT id FROM ceg WHERE kod = 'munkaero');",
        "",
        "-- Partners",
    ]

    # Avoid \\gset — use subquery instead for portability
    lines = [
        "-- Auto-generated: Munkaerő SZSZ seed from Excel 2026-08-31",
        "BEGIN;",
        "",
        f"INSERT INTO ceg (kod, nev) VALUES ({sql_str(CEG_KOD)}, {sql_str('Munkaerő Humánszolgáltató Szociális Szövetkezet')})",
        "ON CONFLICT (kod) DO UPDATE SET nev = EXCLUDED.nev;",
        "",
        "DELETE FROM partner_szerzodes WHERE partner_id IN (SELECT id FROM partner WHERE ceg_id = (SELECT id FROM ceg WHERE kod = 'munkaero'));",
        "DELETE FROM partner_kapcsolattarto WHERE partner_id IN (SELECT id FROM partner WHERE ceg_id = (SELECT id FROM ceg WHERE kod = 'munkaero'));",
        "DELETE FROM projekt WHERE ceg_id = (SELECT id FROM ceg WHERE kod = 'munkaero');",
        "DELETE FROM partner WHERE ceg_id = (SELECT id FROM ceg WHERE kod = 'munkaero');",
        "DELETE FROM munkatars WHERE ceg_id = (SELECT id FROM ceg WHERE kod = 'munkaero');",
        "",
    ]

    for p in partners:
        meta = {
            "kulso_id": p["kulso_id"],
            "cegjegyzekszam": p["cegjegyzekszam"],
            "bankszamlaszam": p["bankszamlaszam"],
            "afakoros": p["afakoros"],
            "nemzetgazdasagi_agazat": p["agazat"],
            "forras": "partner-export-2026-08-31",
        }
        lines.append(
            "INSERT INTO partner (ceg_id, nev, adoszam, cim, iroda, statusz, kapcsolat_tipus, crm_statusz, meta) "
            f"VALUES ((SELECT id FROM ceg WHERE kod = 'munkaero'), {sql_str(p['nev'])}, {sql_str(p['adoszam'])}, "
            f"{sql_str(p['cim'])}, {sql_str(IRODA)}, 'aktív', 'megbízó', 'Megbízó', {sql_json(meta)});"
        )

    lines.append("")
    lines.append("-- Projektek + kapcsolódó szerződések + kapcsolattartók")

    for d in projektek:
        az = str(d.get("Projekt azonosító") or "").strip()
        partner_adoszam = str(d.get("Partner adószáma") or "").strip() or None
        partner_nev = str(d.get("Partner név") or "").strip()
        statusz = norm_status(d.get("Státusz"))
        iroda = str(d.get("Kirendeltség") or IRODA).strip()
        munkatipus = str(d.get("Munkatípus") or "").strip() or None
        nev = f"{az} — {partner_nev}" if partner_nev else az

        meta = {
            "forras": "projekt-adatok-riport-2026-08-31",
            "varmegye": d.get("Vármegye"),
            "partner_iranyitoszam": d.get("Partner - irányítószám"),
            "partner_varos": d.get("Partner - város"),
            "partner_cim": d.get("Partner - cím"),
            "partner_agazat": d.get("Partner - nemzetgazdasági ágazat"),
            "projekt_agazat": d.get("Projekt - ágazat"),
            "projekt_nemzetgazdasagi_agazat": d.get("Projekt - nemzetgazdasági ágazat"),
            "munkatipus": munkatipus,
            "szerzodes_kulso_id": d.get("Kapcsolódó szerződés azonosító"),
            "szerzodes_kelte": parse_date(d.get("Szerződéskötés dátuma")),
            "szerzodes_erv_kezdete": parse_date(d.get("Érvényesség kezdete")),
            "szerzodes_erv_vege": parse_date(d.get("Érvényesség vége")),
            "szerzodes_statusz": d.get("Szerződés státusza"),
            "valtozo_hely": d.get("Változó munkavégzési hely"),
            "munkavegzes_helye": {
                "iranyitoszam": d.get("Munkavégzés helye - irányítószám"),
                "varmegye": d.get("Munkavégzés helye - vármegye"),
                "varos": d.get("Munkavégzés helye - város"),
                "cim": d.get("Munkavégzés helye - cím"),
            },
            "idotartam_tipus": d.get("Projekt időtartam típusa"),
            "projekt_vege": parse_date(d.get("Projekt vége")),
            "leiras": d.get("Leírás"),
            "cimkek": d.get("Címkék"),
            "uzemorvos_figyeles": d.get("Üzemorvosi vizsgálat figyelése beosztás készítésnél"),
            "tudo_figyeles": d.get("Tüdőszűrő vizsgálat figyelése beosztás készítésnél"),
        }

        lines.append(
            "INSERT INTO projekt (ceg_id, azonosito, nev, partner_nev, partner_id, iroda, statusz, meta) VALUES ("
            f"(SELECT id FROM ceg WHERE kod = 'munkaero'), {sql_str(az)}, {sql_str(nev)}, {sql_str(partner_nev)}, "
            f"(SELECT id FROM partner WHERE ceg_id = (SELECT id FROM ceg WHERE kod = 'munkaero') AND adoszam = {sql_str(partner_adoszam)} ORDER BY id LIMIT 1), "
            f"{sql_str(iroda)}, {sql_str(statusz)}, {sql_json(meta)});"
        )

        # Kapcsolattartó
        kt_nev = str(d.get("Elsődleges kapcsolattartó neve") or "").strip()
        if kt_nev and partner_adoszam:
            lines.append(
                "INSERT INTO partner_kapcsolattarto (partner_id, nev, email, mobil, vezetekes, beosztas, szamlazasi, hozzaferes, aktiv) "
                f"SELECT p.id, {sql_str(kt_nev)}, {sql_str(d.get('Email címe'))}, {sql_str(d.get('Telefonszáma - mobil'))}, "
                f"{sql_str(d.get('Telefonszáma - vezetékes'))}, 'Kapcsolattartó', false, 'nincs', true "
                f"FROM partner p WHERE p.ceg_id = (SELECT id FROM ceg WHERE kod = 'munkaero') AND p.adoszam = {sql_str(partner_adoszam)} "
                f"AND NOT EXISTS (SELECT 1 FROM partner_kapcsolattarto k WHERE k.partner_id = p.id AND k.nev = {sql_str(kt_nev)} AND k.email IS NOT DISTINCT FROM {sql_str(d.get('Email címe'))}) "
                "LIMIT 1;"
            )

        # Szerződés a partnerhez (projekt meta is contains link)
        szer_id = d.get("Kapcsolódó szerződés azonosító")
        if szer_id and partner_adoszam:
            tipus = f"projekt-{az}"
            st = str(d.get("Szerződés státusza") or "piszkozat").strip().lower()
            if "érvényes" in st or "ervenyes" in st:
                st_db = "aláírt"
            else:
                st_db = "piszkozat"
            lines.append(
                "INSERT INTO partner_szerzodes (partner_id, tipus, statusz, erv_kezdete, erv_vege, dokumentum_nev) "
                f"SELECT p.id, {sql_str(tipus)}, {sql_str(st_db)}, {sql_str(parse_date(d.get('Érvényesség kezdete')))}::date, "
                f"{sql_str(parse_date(d.get('Érvényesség vége')))}::date, {sql_str(f'Szerződés #{szer_id} / {az}')} "
                f"FROM partner p WHERE p.ceg_id = (SELECT id FROM ceg WHERE kod = 'munkaero') AND p.adoszam = {sql_str(partner_adoszam)} LIMIT 1;"
            )

    lines.append("")
    lines.append("-- Munkatársak + jogosultság mátrix")
    for u in users:
        lines.append(
            "INSERT INTO munkatars (ceg_id, kulso_id, nev, email, szerepkorok, jogosultsag, aktiv) VALUES ("
            f"(SELECT id FROM ceg WHERE kod = 'munkaero'), {u['kulso_id'] if u['kulso_id'] is not None else 'NULL'}, "
            f"{sql_str(u['nev'])}, {sql_str(u['email'])}, {sql_str(u['szerepkorok'])}, {sql_json(u['jogosultsag'])}, true);"
        )

    lines += ["", "COMMIT;", ""]
    OUT_SQL.write_text("\n".join(lines), encoding="utf-8")

    summary = {
        "ceg": CEG_KOD,
        "partners": len(partners),
        "projektek": len(projektek),
        "munkatarsak": len(users),
        "sql": str(OUT_SQL.relative_to(ROOT)),
    }
    OUT_JSON.write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(summary, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
